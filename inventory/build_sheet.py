"""Builds the inventory workbook that seeds the team's Google Sheet.
   in:  inventory/inventory.json (node inventory/build_inventory.mjs)
   out: inventory/DBF-Inventory.xlsx  (upload to Google Drive; it converts to a Google Sheet)
Run: python3 inventory/build_sheet.py
"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
data = json.loads((HERE / 'inventory.json').read_text())
S, rows = data['summary'], data['rows']

F = 'Arial'
BASE = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
HEAD = Font(name=F, size=10, bold=True, color='FFFFFF')
INPUT = Font(name=F, size=10, color='0000FF')
TITLE = Font(name=F, size=14, bold=True)
HFILL = PatternFill('solid', fgColor='16191D')
YELLOW = PatternFill('solid', fgColor='FFF2CC')
THIN = Side(style='thin', color='CBCFD4')
USD = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.00%'

STATUSES = ['To order', 'Ordered', 'Received', 'Inspected OK', 'Rejected', 'To make', 'In progress', 'Made',
            'To source', 'Hold: NWA swap pending', 'Proposed: confirm with NWA', 'Cancelled']
KITS = list(S['kits'].keys())

wb = Workbook()

# ---------------- Read me ----------------
rm = wb.active; rm.title = 'Read me'
lines = [
    ('Design-Build-Fly inventory (DBF-INV-001)', TITLE),
    ('Owner: Tristan Hottier, Inventory Czar. Lift + cruise VTOL, design revision 0.3.1.', BASE),
    ('', BASE),
    ('How to use', BOLD),
    ('Inventory tab: one row per stock line. P- rows are bought, M- rows are made in-house, N- rows are the proposed Nautical Wings carbon props.', BASE),
    ('Edit only the blue cells: Qty ordered, Qty on hand, Status, Location, Order date, Received date, Inspected, Actual USD, Notes.', BASE),
    ('Black cells are formulas (Ext USD, Qty still to get). Status and Inspected use dropdowns; status colours update by themselves.', BASE),
    ('When a delivery arrives: count it against the packing slip, set Qty on hand, Received date and Inspected, label the item with its ID and put it in its kit bin (K1 to K7).', BASE),
    ('Budget tab: totals against the $2,000 cap, the crash reserve, spend by kit and the effect of the Nautical Wings prop swap.', BASE),
    ('', BASE),
    ('Where the numbers come from', BOLD),
    (f"Rows were generated on {S['generated']} from docs/purchasing/C0_BOM.csv (prices checked {S['bomPricesChecked']}) and the CAD assembly in cad/params.txt, by inventory/build_inventory.mjs in the DBF-VTOL repo.", BASE),
    ('Nothing has been ordered or built yet, so every on-hand quantity starts at 0. Re-check prices before ordering.', BASE),
    ('Overview page: Inventory-Overview.html in the DBF-VTOL repo.', BASE),
]
for i, (t, f) in enumerate(lines, 1):
    c = rm.cell(row=i, column=1, value=t); c.font = f; c.alignment = Alignment(wrap_text=True, vertical='top')
rm.column_dimensions['A'].width = 120

# ---------------- Inventory ----------------
inv = wb.create_sheet('Inventory')
cols = ['ID', 'Item', 'Type', 'Category', 'Kit', 'CAD part', 'Qty needed', 'Qty ordered', 'Qty on hand',
        'Qty still to get', 'Status', 'Vendor', 'Unit USD', 'Ext USD', 'Actual USD', 'Made from', 'Process',
        'Location', 'Order date', 'Received date', 'Inspected', 'Source', 'Notes']
C = {name: get_column_letter(i + 1) for i, name in enumerate(cols)}
EDIT = {'Qty ordered', 'Qty on hand', 'Status', 'Location', 'Order date', 'Received date', 'Inspected', 'Actual USD', 'Notes'}
for j, name in enumerate(cols, 1):
    c = inv.cell(row=1, column=j, value=name); c.font = HEAD; c.fill = HFILL
    c.alignment = Alignment(vertical='center', wrap_text=True)
inv.cell(row=1, column=1).comment = Comment(
    f"Source: DBF-VTOL repo, inventory/inventory.json (from docs/purchasing/C0_BOM.csv, prices checked {S['bomPricesChecked']}, "
    f"and cad/params.txt), generated {S['generated']} by inventory/build_inventory.mjs.", 'Claude')

for i, r in enumerate(rows, 2):
    for j, name in enumerate(cols, 1):
        if name == 'Qty still to get':
            v = f'=MAX(0,{C["Qty needed"]}{i}-{C["Qty on hand"]}{i}-N({C["Qty ordered"]}{i}))'
        elif name == 'Ext USD':
            v = f'={C["Qty needed"]}{i}*{C["Unit USD"]}{i}'
        elif name == 'Actual USD':
            v = None
        else:
            v = r.get(name, '')
            if v == '' or v is None: v = None
            if name in ('Qty ordered',) and r['Type'] == 'Make': v = None
        c = inv.cell(row=i, column=j, value=v); c.font = INPUT if name in EDIT else BASE
        c.border = Border(bottom=THIN)
        if name in ('Unit USD', 'Ext USD', 'Actual USD'): c.number_format = USD
        if name in ('Order date', 'Received date'): c.number_format = 'yyyy-mm-dd'
        if name in ('Item', 'Process', 'Notes', 'Source'): c.alignment = Alignment(wrap_text=True, vertical='top')
        else: c.alignment = Alignment(vertical='top')
last = len(rows) + 1
widths = {'ID': 7, 'Item': 46, 'Type': 7, 'Category': 16, 'Kit': 5, 'CAD part': 16, 'Qty needed': 8, 'Qty ordered': 8,
          'Qty on hand': 8, 'Qty still to get': 9, 'Status': 24, 'Vendor': 18, 'Unit USD': 10, 'Ext USD': 10,
          'Actual USD': 10, 'Made from': 9, 'Process': 30, 'Location': 10, 'Order date': 11, 'Received date': 11,
          'Inspected': 9, 'Source': 30, 'Notes': 40}
for name, w in widths.items(): inv.column_dimensions[C[name]].width = w
inv.row_dimensions[1].height = 30
inv.freeze_panes = 'C2'
inv.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{last}'

# dropdowns (to row 500 so new rows get them)
lst = wb.create_sheet('Lists')
lst['A1'] = 'Status'; lst['B1'] = 'Kit'; lst['C1'] = 'Kit name'; lst['D1'] = 'Inspected'; lst['E1'] = 'Type'
for c in ('A1', 'B1', 'C1', 'D1', 'E1'): lst[c].font = BOLD
for i, s in enumerate(STATUSES, 2): lst.cell(row=i, column=1, value=s).font = BASE
for i, k in enumerate(KITS, 2):
    lst.cell(row=i, column=2, value=k).font = BASE; lst.cell(row=i, column=3, value=S['kits'][k]).font = BASE
for i, s in enumerate(['Y', 'N', 'Partial'], 2): lst.cell(row=i, column=4, value=s).font = BASE
for i, s in enumerate(['Buy', 'Make', 'NWA'], 2): lst.cell(row=i, column=5, value=s).font = BASE
lst.column_dimensions['A'].width = 28; lst.column_dimensions['C'].width = 26
N = 500
for col, rng in (('Status', f'Lists!$A$2:$A${len(STATUSES)+1}'), ('Kit', f'Lists!$B$2:$B${len(KITS)+1}'),
                 ('Inspected', 'Lists!$D$2:$D$4'), ('Type', 'Lists!$E$2:$E$4')):
    dv = DataValidation(type='list', formula1='=' + rng, allow_blank=True, showErrorMessage=False)
    inv.add_data_validation(dv); dv.add(f'{C[col]}2:{C[col]}{N}')

# status colours (whole row range of the Status cell)
st = C['Status']
for text, colour in (('Inspected OK', 'D9EAD3'), ('Received', 'D9EAD3'), ('Made', 'D9EAD3'), ('Rejected', 'F4CCCC'),
                     ('To source', 'F4CCCC'), ('To order', 'FCE5CD'), ('Ordered', 'FFF2CC'), ('In progress', 'FFF2CC'),
                     ('Hold', 'D9D2E9'), ('Proposed', 'D9D2E9'), ('To make', 'CFE2F3')):
    inv.conditional_formatting.add(f'{st}2:{st}{N}', FormulaRule(formula=[f'LEFT(${st}2,{len(text)})="{text}"'],
                                                                   fill=PatternFill('solid', fgColor=colour)))

# ---------------- Budget ----------------
b = wb.create_sheet('Budget', 1)
b['A1'] = 'Budget against the $2,000 cap'; b['A1'].font = TITLE
b['A2'] = 'Blue cells are inputs; black cells are formulas over the Inventory tab.'; b['A2'].font = BASE
inputs = [
    (4, 'Hardware budget cap', S['cap'], USD, 'Source: docs/DESIGN-BASIS.md R6 (fixed, includes tax and shipping)'),
    (5, 'Sales/use tax rate (Berkeley)', S['taxRate'], PCT, 'Source: C0_BOM.csv, CDTFA rate 2026-10-01'),
    (6, 'Shipping allowance (estimate)', S['shipping'], USD, 'Source: C0_BOM.csv; estimate, price at checkout'),
]
b['A3'] = 'Inputs'; b['A3'].font = BOLD
for r, label, v, fmt, src in inputs:
    b.cell(row=r, column=1, value=label).font = BASE
    c = b.cell(row=r, column=2, value=v); c.font = INPUT; c.number_format = fmt; c.comment = Comment(src, 'Claude')
    b.cell(row=r, column=3, value=src).font = Font(name=F, size=9, color='6C737C')
I = 'Inventory'
ext, typ, status, kitc, act = (f"{I}!${C['Ext USD']}$2:${C['Ext USD']}${N}", f"{I}!${C['Type']}$2:${C['Type']}${N}",
                               f"{I}!${C['Status']}$2:${C['Status']}${N}", f"{I}!${C['Kit']}$2:${C['Kit']}${N}",
                               f"{I}!${C['Actual USD']}$2:${C['Actual USD']}${N}")
b['A8'] = 'Plan (BOM prices)'; b['A8'].font = BOLD
calc = [
    (9, 'Goods subtotal (Buy lines)', f'=SUMIFS({ext},{typ},"Buy")'),
    (10, 'Sales/use tax', '=B9*B5'),
    (11, 'Shipping allowance', '=B6'),
    (12, 'Committed spend', '=B9+B10+B11'),
    (13, 'Crash reserve left under the cap', '=B4-B12'),
    (14, 'Reserve check: BOM says $102.08', '=IF(ROUND(B13,1)=102.1,"matches C0_BOM.csv","differs from C0_BOM.csv: re-check")'),
]
for r, label, f in calc:
    b.cell(row=r, column=1, value=label).font = BOLD if r in (12, 13) else BASE
    c = b.cell(row=r, column=2, value=f); c.font = BOLD if r in (12, 13) else BASE
    if r != 14: c.number_format = USD
b['A16'] = 'Actual spend (fill Actual USD on the Inventory tab, tax and shipping included)'; b['A16'].font = BOLD
b['A17'] = 'Spent so far'; b['B17'] = f'=SUM({act})'; b['B17'].number_format = USD
b['A18'] = 'Left under the cap'; b['B18'] = '=B4-B17'; b['B18'].number_format = USD
for c in ('A17', 'B17', 'A18', 'B18'): b[c].font = BASE

b['A20'] = 'Nautical Wings carbon props (proposed, free)'; b['A20'].font = BOLD
b['A21'] = 'Retail prop lines on hold (goods)'; b['B21'] = f'=SUMIFS({ext},{status},"Hold*")'
b['A22'] = 'Same with tax'; b['B22'] = '=B21*(1+B5)'
b['A23'] = 'Reserve if the NWA props are confirmed (EST)'; b['B23'] = '=B13+B22'
b['C23'] = 'EST: assumes the props are free with no shipping or import cost; shipping allowance unchanged'
b['C23'].font = Font(name=F, size=9, color='6C737C')
for r in (21, 22, 23):
    b.cell(row=r, column=1).font = BASE; b.cell(row=r, column=2).font = BASE; b.cell(row=r, column=2).number_format = USD

b['A25'] = 'By kit'; b['A25'].font = BOLD
hdr = ['Kit', 'Name', 'Lines', 'Still to get (lines)', 'Plan goods (USD)']
for j, h in enumerate(hdr, 1):
    c = b.cell(row=26, column=j, value=h); c.font = HEAD; c.fill = HFILL
for i, k in enumerate(KITS, 27):
    lr = i - 25  # Lists row for this kit
    b.cell(row=i, column=1, value=f'=Lists!B{lr}').font = BASE
    b.cell(row=i, column=2, value=f'=Lists!C{lr}').font = BASE
    b.cell(row=i, column=3, value=f'=COUNTIFS({kitc},A{i})').font = BASE
    b.cell(row=i, column=4, value=f"=COUNTIFS({kitc},A{i},{I}!${C['Qty still to get']}$2:${C['Qty still to get']}${N},\">0\")").font = BASE
    c = b.cell(row=i, column=5, value=f'=SUMIFS({ext},{kitc},A{i},{typ},"Buy")'); c.font = BASE; c.number_format = USD
t = 27 + len(KITS)
b.cell(row=t, column=1, value='Total').font = BOLD
for col in ('C', 'D', 'E'):
    c = b[f'{col}{t}']; c.value = f'=SUM({col}27:{col}{t-1})'; c.font = BOLD
b[f'E{t}'].number_format = USD
b.column_dimensions['A'].width = 44; b.column_dimensions['B'].width = 22; b.column_dimensions['C'].width = 14
b.column_dimensions['D'].width = 18; b.column_dimensions['E'].width = 18

# ---------------- Data Sources ----------------
ds = wb.create_sheet('Data Sources')
h = ['Written to', 'Source', 'Object', 'Obtained via', 'Run at']
for j, x in enumerate(h, 1):
    c = ds.cell(row=1, column=j, value=x); c.font = HEAD; c.fill = HFILL
ds.append([f'Inventory!A1:{get_column_letter(len(cols))}{last}', 'DBF-VTOL repo', 'inventory/inventory.json (C0_BOM.csv + cad/params.txt)',
           'Generated by inventory/build_inventory.mjs, run by Claude', S['generated']])
for c in ds[2]: c.font = BASE
for col, w in zip('ABCDE', (22, 16, 52, 48, 12)): ds.column_dimensions[col].width = w

wb._sheets = [wb[n] for n in ('Read me', 'Budget', 'Inventory', 'Lists', 'Data Sources')]
out = HERE / 'DBF-Inventory.xlsx'
wb.save(out)
print('wrote', out, out.stat().st_size, 'bytes,', len(rows), 'rows')

# ---------------- single-tab CSV for the Google Sheet (Drive converts it; formulas stay live) ----------------
import csv, io
gcols = ['ID', 'Item', 'Type', 'Kit', 'CAD part', 'Qty needed', 'Qty ordered', 'Qty on hand', 'Qty still to get', 'Status',
         'Vendor', 'Unit USD', 'Ext USD', 'Actual USD', 'Made from', 'Process', 'Location', 'Order date', 'Received date',
         'Inspected', 'Notes']
G = {name: get_column_letter(i + 1) for i, name in enumerate(gcols)}
HR = 15                      # header row
D0, DN = HR + 1, 1000        # data rows (open to 1000 so new rows count)
rng = lambda c: f'${G[c]}${D0}:${G[c]}${DN}'
grid = [[''] * len(gcols) for _ in range(HR - 1)]
def put(r, c, v): grid[r - 1][c - 1] = v
put(1, 1, 'Design-Build-Fly inventory (DBF-INV-001)')
put(1, 2, f"Owner: Tristan Hottier, Inventory Czar. Rows generated {S['generated']} from C0_BOM.csv (prices checked {S['bomPricesChecked']}) and cad/params.txt. Edit Qty ordered, Qty on hand, Status, Location, dates, Inspected, Actual USD, Notes.")
put(2, 1, 'Inputs'); put(2, 3, 'Kit name'); put(2, 4, 'Kit'); put(2, 5, 'Lines'); put(2, 6, 'Still to get'); put(2, 7, 'Plan goods USD')
put(3, 1, 'Budget cap (DESIGN-BASIS.md R6)'); put(3, 2, S['cap'])
put(4, 1, 'Tax rate (C0_BOM.csv, CDTFA 2026-10-01)'); put(4, 2, S['taxRate'])
put(5, 1, 'Shipping allowance (estimate)'); put(5, 2, S['shipping'])
put(6, 1, 'Goods subtotal'); put(6, 2, f'=SUMIFS({rng("Ext USD")},{rng("Type")},"Buy")')
put(7, 1, 'Tax'); put(7, 2, '=B6*B4')
put(8, 1, 'Committed spend'); put(8, 2, '=B6+B7+B5')
put(9, 1, 'Crash reserve'); put(9, 2, '=B3-B8')
put(10, 1, 'Spent so far (Actual USD)'); put(10, 2, f'=SUM({rng("Actual USD")})')
put(11, 1, 'NWA swap: retail props on hold, with tax'); put(11, 2, f'=SUMIFS({rng("Ext USD")},{rng("Status")},"Hold*")*(1+B4)')
put(12, 1, 'Reserve if NWA props confirmed (EST: free, no shipping)'); put(12, 2, '=B9+B11')
for i, k in enumerate(KITS):
    r = 3 + i
    put(r, 4, k); put(r, 3, S['kits'][k])
    put(r, 5, f'=COUNTIFS({rng("Kit")},D{r})')
    put(r, 6, f'=COUNTIFS({rng("Kit")},D{r},{rng("Qty still to get")},">0")')
    put(r, 7, f'=SUMIFS({rng("Ext USD")},{rng("Kit")},D{r},{rng("Type")},"Buy")')
grid.append(gcols)
for i, r in enumerate(rows):
    n = D0 + i; line = []
    for c in gcols:
        if c == 'Qty still to get': v = f'=MAX(0,{G["Qty needed"]}{n}-{G["Qty on hand"]}{n}-N({G["Qty ordered"]}{n}))'
        elif c == 'Ext USD': v = f'={G["Qty needed"]}{n}*{G["Unit USD"]}{n}'
        elif c == 'Actual USD': v = ''
        elif c == 'Qty ordered' and r['Type'] == 'Make': v = ''
        else: v = r.get(c, '')
        line.append('' if v is None else v)
    grid.append(line)
buf = io.StringIO(); csv.writer(buf, lineterminator='\n').writerows(grid)
(HERE / 'DBF-Inventory-sheet.csv').write_text(buf.getvalue())
print('wrote DBF-Inventory-sheet.csv', len(buf.getvalue()), 'chars')
